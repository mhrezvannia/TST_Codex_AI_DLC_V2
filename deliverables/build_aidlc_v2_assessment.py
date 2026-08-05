from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\TST_Codex_integ")
OUT = ROOT / "deliverables" / "AI-DLC_v2_Assessment_Booking_Quote_to_Cash.docx"
REFERENCE_IMAGE = ROOT / "design-inputs" / "claude-ui-export" / "screenshots" / "platform-bundled.png"

NAVY = "0B2545"
BLUE = "2E74B5"
TEAL = "136B45"
AMBER = "8A5200"
RED = "B42318"
INK = "17212B"
MUTED = "64717D"
LIGHT = "F2F4F7"
BLUE_LIGHT = "E8EEF5"
TEAL_LIGHT = "E7F4EE"
AMBER_LIGHT = "FBF0DC"
RED_LIGHT = "FBE9E7"
WHITE = "FFFFFF"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:tblHeader")
    flag.set(qn("w:val"), "true")
    tr_pr.append(flag)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    flag = OxmlElement("w:cantSplit")
    tr_pr.append(flag)


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            width = widths_dxa[min(index, len(widths_dxa) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def style_table(table, widths_dxa, header=True, font_size=9.2):
    set_table_geometry(table, widths_dxa)
    if header:
        set_repeat_header(table.rows[0])
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            if header and r_idx == 0:
                set_cell_shading(cell, NAVY)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, LIGHT)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    set_run_font(run, size=font_size, color=WHITE if header and r_idx == 0 else INK, bold=header and r_idx == 0)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for node in (fld_begin, instr, fld_sep, text, fld_end):
        run._r.append(node)


def create_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abs = max(abstract_ids, default=0) + 1
    next_num = max(num_ids, default=0) + 1

    def add_definition(abstract_id, num_id, fmt, text, left, hanging):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abstract_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(left))
        tabs.append(tab)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), str(hanging))
        ppr.extend([tabs, ind])
        lvl.extend([start, num_fmt, lvl_text, suff, ppr])
        if fmt == "bullet":
            rpr = OxmlElement("w:rPr")
            rfonts = OxmlElement("w:rFonts")
            rfonts.set(qn("w:ascii"), "Arial")
            rfonts.set(qn("w:hAnsi"), "Arial")
            rpr.append(rfonts)
            lvl.append(rpr)
        abstract.append(lvl)
        numbering.append(abstract)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abs_id = OxmlElement("w:abstractNumId")
        abs_id.set(qn("w:val"), str(abstract_id))
        num.append(abs_id)
        numbering.append(num)

    add_definition(next_abs, next_num, "bullet", "•", 720, 360)
    add_definition(next_abs + 1, next_num + 1, "decimal", "%1.", 720, 360)
    return next_num, next_num + 1


def apply_num(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    nid = OxmlElement("w:numId")
    nid.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, nid])


def add_bullet(doc, text, bullet_num, bold_prefix=None):
    p = doc.add_paragraph()
    apply_num(p, bullet_num)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_number(doc, text, decimal_num):
    p = doc.add_paragraph()
    apply_num(p, decimal_num)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    set_run_font(p.add_run(text))
    return p


def add_callout(doc, label, text, fill=BLUE_LIGHT, accent=BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.12
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), accent)
    borders.append(left)
    ppr.append(borders)
    r = p.add_run(f"{label}: ")
    set_run_font(r, bold=True, color=accent)
    r2 = p.add_run(text)
    set_run_font(r2, color=INK)
    return p


def add_source_line(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, name="Consolas", size=8.3, color=MUTED)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    return p


def add_page_break(doc):
    doc.add_page_break()


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.78)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.7)
normal.font.color.rgb = RGBColor.from_string(INK)
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for style_name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, NAVY, 8, 4),
):
    style = styles[style_name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

bullet_num, decimal_num = create_numbering(doc)

# Running header and footer: quiet, no decorative rule.
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hp.paragraph_format.space_after = Pt(0)
set_run_font(hp.add_run("AI-DLC v2 Research Brief"), size=8.5, color=MUTED, bold=True)
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_run_font(fp.add_run("Internal assessment  |  "), size=8.5, color=MUTED)
add_field(fp, "PAGE")

# Cover / memo masthead.
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.paragraph_format.space_after = Pt(4)
set_run_font(p.add_run("RESEARCH AND DECISION BRIEF"), size=10, color=BLUE, bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(6)
set_run_font(p.add_run("AI-DLC v2: What It Produces, Where It Helps, and Where It Does Not"), size=24, color=NAVY, bold=True)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
set_run_font(p.add_run("Evidence-based assessment using W1-01 Booking Quote-to-Cash"), size=13, color=MUTED)

meta_data = [
    ("Prepared for", "Manager and CTO"),
    ("Decision", "Whether and how to use AI-DLC v2"),
    ("Case study", "W1-01 Booking Quote-to-Cash"),
    ("Evidence date", "19 July 2026"),
    ("Evidence basis", "Repository artifacts, current frontend source, test/deployment records, and committed visual references"),
]
for label, value in meta_data:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.05
    set_run_font(p.add_run(f"{label}: "), size=9.5, color=NAVY, bold=True)
    set_run_font(p.add_run(value), size=9.5, color=INK)

doc.add_paragraph()
add_callout(
    doc,
    "Bottom line",
    "AI-DLC v2 is useful as a traceable engineering and delivery framework, but it is not an automatic requirements-to-exact-product compiler. For UI-critical work, it must be paired with an approved design authority (normally a versioned Figma library and feature prototype), automated visual checks, and a release gate that cannot be marked complete while live evidence is blocked.",
    fill=TEAL_LIGHT,
    accent=TEAL,
)

add_heading(doc, "Executive conclusion", 1)
p = doc.add_paragraph()
p.add_run("Recommendation: ").bold = True
p.add_run("adopt AI-DLC v2 conditionally, not as a standalone guarantee of fidelity.")
p = doc.add_paragraph("The Booking Quote-to-Cash record demonstrates both sides of the methodology. It created unusually strong traceability from intent through requirements, stories, design notes, code summaries, tests, deployment, observability, and drift reporting. It also demonstrates that documentation and stage completion do not force the implemented UI to match the approved design, and do not prove the live outcome when deployment is blocked.")

table = doc.add_table(rows=1, cols=3)
for cell, text in zip(table.rows[0].cells, ["Dimension", "Case-study result", "Management reading"]):
    cell.text = text
rows = [
    ("Business scope and traceability", "Strong", "The thin commercial journey is explicit and decomposed across real service seams."),
    ("Functional implementation", "Substantial", "Routes, create, validation, pricing, confirmation, and status polling exist in source and deterministic tests pass."),
    ("UX/design artifacts", "Detailed but text-based", "ASCII wireframes and Markdown interaction specifications exist; no native Figma artifact exists."),
    ("UI implementation fidelity", "Weak", "Current Booking source contradicts several approved visual/component rules."),
    ("Live release proof", "Blocked", "Fresh full-stack acceptance and performance workloads did not run to completion."),
    ("Overall decision", "Use with controls", "Keep AI-DLC for engineering rigor; add design and evidence gates for exact output."),
]
for item in rows:
    cells = table.add_row().cells
    for cell, text in zip(cells, item):
        cell.text = text
style_table(table, [2200, 2000, 5160])

add_heading(doc, "What AI-DLC v2 actually outputs", 1)
p = doc.add_paragraph("AI-DLC v2 does not have one single output. Its output is a chain of decision artifacts, implementation artifacts, and evidence. In this repository, the completed W1 record contains 32 lifecycle stages and the following classes of deliverables.")

table = doc.add_table(rows=1, cols=4)
for cell, text in zip(table.rows[0].cells, ["Lifecycle phase", "Typical output", "W1 examples", "What it proves"]):
    cell.text = text
phase_rows = [
    ("Ideation", "Intent, research, feasibility, scope, team, rough UX, approval", "intent-statement.md; scope-document.md; wireframes.md; user-flow.md; initiative-brief.md", "What problem and thin outcome are approved; early UX direction"),
    ("Inception", "Requirements, stories, refined UX, application design, units, delivery plan", "requirements.md; stories.md; mockups.md; interaction-spec.md; components.md; unit-of-work.md; bolt-plan.md", "A buildable and traceable design, including error states and ownership"),
    ("Construction", "Per-unit functional/NFR/infrastructure design, code, tests, CI", "business-rules.md; frontend-components.md; code-summary.md; build-test-results.md; quality-gates.md", "What was built and what deterministic checks passed"),
    ("Operation", "Deployment, health, observability, incident, performance, feedback", "deployment-log.md; health-check-report.md; dashboards.md; alarms.md; load-test-results.md; drift-report.md", "What ran live, what remains blocked, and what must be corrected"),
]
for item in phase_rows:
    cells = table.add_row().cells
    for cell, text in zip(cells, item):
        cell.text = text
style_table(table, [1300, 2600, 3000, 2460], font_size=8.7)

add_callout(doc, "Important distinction", "A completed artifact chain proves process traceability. It does not automatically prove product correctness, visual fidelity, usability, production readiness, or stakeholder acceptance.", fill=AMBER_LIGHT, accent=AMBER)

add_heading(doc, "What it does not automatically produce", 2)
for text in [
    "A guaranteed pixel-accurate UI or a native Figma source file.",
    "A correct design system unless one is supplied, approved, versioned, and enforced.",
    "A reliable business decision when inputs are ambiguous or stakeholders disagree.",
    "A passed live release when infrastructure, credentials, registries, or test environments are unavailable.",
    "Independent acceptance: AI-generated review verdicts can still be self-referential or incomplete.",
    "Protection from scope mismatch when stakeholders expect features that the intent explicitly defers.",
]:
    add_bullet(doc, text, bullet_num)

add_page_break(doc)
add_heading(doc, "Pros of AI-DLC v2", 1)
pros = [
    ("End-to-end traceability", "The method leaves an auditable path from intent through requirements, stories, designs, code, tests, and runtime evidence."),
    ("Thin vertical delivery", "It discourages isolated backend work and makes the user outcome cross UI, API, domain, persistence, messaging, and read models."),
    ("Better enterprise seam discipline", "Contracts, ownership, idempotency, schemas, and service boundaries become explicit rather than implied."),
    ("State and failure coverage", "Loading, empty, unavailable, validation, pending, retry, duplicate, stale-event, and restart behavior are designed before implementation."),
    ("Non-functional requirements", "Latency, accessibility, security, reliability, observability, and deployment are part of the record instead of late checklists."),
    ("Review and correction loop", "The W1 code review found real authorization, migration, filter-state, and field-level error issues, and later iterations corrected them."),
    ("Honest evidence retention", "Blocked deployment and unexecuted performance tests were preserved rather than relabeled as passed."),
    ("Brownfield suitability", "The method is strong where legacy compatibility, multiple services, migrations, contracts, and operational constraints matter."),
]
for title, body in pros:
    p = doc.add_paragraph()
    apply_num(p, decimal_num)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(f"{title}. ")
    set_run_font(r, bold=True, color=NAVY)
    set_run_font(p.add_run(body))

add_heading(doc, "Cons and failure modes", 1)
cons = [
    ("Artifact volume", "A full lifecycle can create many documents. Without disciplined reading and automated trace checks, documentation becomes expensive inventory."),
    ("False confidence from completion", "The W1 state says 32/32 stages complete, while the release evidence remains blocked. Teams can confuse process completion with outcome completion."),
    ("Text-first UI weakness", "ASCII layouts and prose describe structure, but cannot lock spacing, typography, component anatomy, responsive composition, or interaction feel."),
    ("Self-consistency is not external correctness", "AI can produce artifacts that agree with each other and still miss the product stakeholder's unstated expectation."),
    ("Gate weakness", "If a requirement is not converted into an executable or human approval gate, implementation may diverge without blocking merge."),
    ("Review independence", "The refined mockup record explicitly says its final product verdict was not independent after two reviewer processes timed out."),
    ("Environment dependency", "Live evidence can fail because of registry, proxy, image, credential, or browser-tool availability even when source tests pass."),
    ("Late design-system work", "Deferring a design system may be correct for a thin feature, but it predictably creates visual inconsistency and later migration work."),
]
for title, body in cons:
    p = doc.add_paragraph()
    apply_num(p, decimal_num)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(f"{title}. ")
    set_run_font(r, bold=True, color=RED)
    set_run_font(p.add_run(body))

add_heading(doc, "The CTO's Figma-first proposal", 1)
add_callout(doc, "Assessment", "The CTO is right that visual design needs an authoritative artifact before UI construction. The strongest sequence is not 'Figma before understanding the problem'; it is 'minimal intent and journey first, then Figma design authority before detailed UI acceptance and code.'", fill=BLUE_LIGHT, accent=BLUE)

table = doc.add_table(rows=1, cols=3)
for cell, text in zip(table.rows[0].cells, ["Approach", "Benefit", "Risk"]):
    cell.text = text
comparison = [
    ("Figma before any intent or workflow", "Early visual alignment", "May beautify the wrong workflow, omit domain constraints, and create rework."),
    ("Stories before design authority", "Fast backlog creation", "Stories encode vague UI language; AI fills the visual gaps with its own defaults."),
    ("Recommended dual track", "Business and visual intent converge before code", "Requires product, design, and engineering to agree on shared gates."),
]
for item in comparison:
    cells = table.add_row().cells
    for cell, text in zip(cells, item):
        cell.text = text
style_table(table, [2200, 3300, 3860])

add_heading(doc, "Recommended order for UI-heavy work", 2)
for text in [
    "Approve the business intent, actor, journey, domain boundaries, data constraints, and measurable Definition of Done.",
    "Create or approve the Figma design-system library: tokens, typography, grid, primitives, states, accessibility rules, and responsive behavior.",
    "Create feature wireframes/prototypes for the real journey, including loading, empty, error, permission, pending, and recovery states.",
    "Write detailed user stories and acceptance criteria that reference immutable Figma frame/component IDs and versions.",
    "Use AI-DLC for application design, units, code generation, testing, security, deployment, and operations, while enforcing the approved design artifacts.",
    "Block merge on visual regression, accessibility, design-token/lint checks, and named product/design approval.",
    "Block release completion until the live acceptance manifest is PASS; preserve BLOCKED as incomplete rather than complete-with-notes.",
]:
    add_number(doc, text, decimal_num)

add_page_break(doc)
add_heading(doc, "Case study: W1-01 Booking Quote-to-Cash", 1)
p = doc.add_paragraph("The original intent was a thin carrier-side commercial spine: create one booking, validate live reference data, obtain a real Charge quote, confirm through a real Kafka event, open a CMM journey, return a movement status, and render it on Booking detail. It explicitly deferred app shell/auth and the full design-system foundation.")
add_source_line(doc, "Source: docs/examples/booking-quote-to-cash/intent-statement.md; aidlc/.../260714-booking-quote-cash/ideation/intent-capture/intent-statement.md")

add_heading(doc, "Wireframe-like and design files found", 2)
table = doc.add_table(rows=1, cols=4)
for cell, text in zip(table.rows[0].cells, ["Artifact", "Format", "What it contains", "Authority level"]):
    cell.text = text
artifact_rows = [
    ("design-inputs/claude-ui-export/Booking Directions.dc.html", "Interactive/static HTML", "Three booking visual directions; Direction B is the operational-console reference", "Visual reference, not Figma"),
    ("design-inputs/claude-ui-export/LinerCore Booking Directions.html", "Bundled HTML", "Packaged visual-directions export", "Visual reference"),
    ("design-inputs/claude-ui-export/LinerCore Platform.html and LinerCore.dc.html", "HTML", "Broader platform shell and cross-module journey", "Broader-than-W1 concept"),
    ("design-inputs/.../screenshots/bundled-check.png", "PNG", "Captured booking visual-direction board", "Static visual evidence"),
    ("design-inputs/.../screenshots/platform-bundled.png", "PNG", "Captured platform/booking concept", "Static visual evidence"),
    ("ideation/rough-mockups/wireframes.md", "Markdown + ASCII", "List, create, detail, states, responsive notes, components", "Approved rough UX"),
    ("ideation/rough-mockups/user-flow.md", "Markdown", "Happy path, recovery, lifecycle, navigation", "Approved behavior"),
    ("inception/refined-mockups/mockups.md", "Markdown + ASCII", "Refined list/create/detail composition and states", "Approved structural design"),
    ("inception/refined-mockups/interaction-spec.md", "Markdown", "Routes, components, state behavior, accessibility", "Approved interaction contract"),
    ("inception/refined-mockups/design-system-mapping.md", "Markdown", "Token and @erp/ui primitive mapping", "Approved implementation rule"),
    ("inception/refined-mockups/accessibility-checklist.md", "Markdown checklist", "WCAG, keyboard, responsive, verification matrix", "Acceptance checklist; boxes remain unchecked in this artifact"),
]
for item in artifact_rows:
    cells = table.add_row().cells
    for cell, text in zip(cells, item):
        cell.text = text
style_table(table, [2600, 1400, 3460, 1900], font_size=8.2)

add_callout(doc, "Figma finding", "No .fig, .sketch, or .xd file was found for W1-01, and no Figma reference appears in the W1 or design-input artifacts. The AI-DLC wireframes and refined mockups are text/ASCII artifacts, while the strongest visual input is an HTML/PNG concept export.", fill=RED_LIGHT, accent=RED)

if REFERENCE_IMAGE.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    inline = run.add_picture(str(REFERENCE_IMAGE), width=Inches(6.35))
    inline._inline.docPr.set("descr", "Committed LinerCore platform and Booking visual concept used as an intended visual reference")
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)
    set_run_font(caption.add_run("Figure 1. Committed platform/Booking visual reference. It includes broader shell elements explicitly outside the W1 thin-slice scope."), size=8.5, color=MUTED, italic=True)

add_heading(doc, "Intent-to-output comparison", 1)
table = doc.add_table(rows=1, cols=5)
for cell, text in zip(table.rows[0].cells, ["Intent element", "Artifact expectation", "Implemented evidence", "Result", "Comment"]):
    cell.text = text
compare_rows = [
    ("Stable list/create/detail routes", "/bookings, /bookings/new, /bookings/{id}", "All three route files exist; production build emitted routes", "Mostly aligned", "Basic navigation and state preservation exist"),
    ("Contract-true thin booking", "One leg, one equipment, USD/FCL dry", "Create form submits routing[], equipment[], currency and cargo mode", "Aligned", "Canonical shape is present"),
    ("Live reference validation", "Server-backed canonical references and blocked invalid states", "Reference options, validate command, field errors, retry behavior", "Substantially aligned", "Input implementation differs from approved Combobox primitive"),
    ("Real Charge pricing", "Persisted itemized quote/basis; manual state on failure", "Pricing command and snapshot/manual sections exist; service tests pass", "Substantially aligned", "Fresh live full-stack proof remains blocked"),
    ("Confirm and CMM status", "Async event, pending state, returned status on detail", "Confirm action and 30-second status polling exist; backend/contract tests pass", "Substantially aligned", "Current live journey not proven by final run"),
    ("Operational detail composition", "Lifecycle strip, six tabs, 320px quote rail, confirm dialog", "Continuous detail grid, pricing section, journey section, lifecycle event list", "Low alignment", "Major approved composition is absent"),
    ("Design-system inheritance", "IBM Plex, @erp/ui tokens/primitives, no local hardcoded colors", "Inter and hardcoded local CSS; no @erp/ui import/token usage in Booking source", "Not aligned", "Direct contradiction of design-system-mapping.md"),
    ("Responsive/accessibility proof", "360/768/1200/1440 no-overlap, keyboard, WCAG AA", "One 720px CSS breakpoint; selected component tests; no committed W1 screenshot proof", "Partially aligned", "Checklist is detailed but not visibly signed off"),
    ("Observed live Definition of Done", "Fresh full stack, nginx route, audits, replay, measured SLOs", "Deployment blocked at image pull; nginx probe failed; performance not executed", "Not proven", "Correctly recorded as BLOCKED in evidence"),
]
for item in compare_rows:
    cells = table.add_row().cells
    for cell, text in zip(cells, item):
        cell.text = text
style_table(table, [1500, 2150, 2400, 1200, 2110], font_size=7.7)

add_heading(doc, "How similar is the implementation?", 2)
p = doc.add_paragraph("A single similarity percentage would hide the most important finding: functional fidelity and visual fidelity are very different. The following provisional score is a transparent source-based rubric, not a pixel-diff result. A fresh browser could not be connected in this review session, and no committed W1 implementation screenshot was found.")

score = doc.add_table(rows=1, cols=5)
for cell, text in zip(score.rows[0].cells, ["Dimension", "Weight", "Score", "Weighted", "Evidence reading"]):
    cell.text = text
score_rows = [
    ("Business and functional intent", "40%", "4.0 / 5", "32 / 40", "The thin journey is substantially represented in code/tests."),
    ("UX behavior and states", "25%", "2.5 / 5", "12.5 / 25", "Many states exist, but approved navigation/composition is incomplete."),
    ("Visual/design-system fidelity", "20%", "1.0 / 5", "4 / 20", "Font, tokens, primitives, tabs, quote rail, and dialog do not match."),
    ("Live acceptance evidence", "15%", "1.5 / 5", "4.5 / 15", "Partial stack health exists; final fresh release proof is blocked."),
    ("Provisional overall", "100%", "-", "53 / 100", "Useful functional slice, not an exact UI or fully proven release."),
]
for item in score_rows:
    cells = score.add_row().cells
    for cell, text in zip(cells, item):
        cell.text = text
style_table(score, [2500, 1000, 1100, 1100, 3660], font_size=8.5)

add_callout(doc, "UI conclusion", "The implemented Booking UI is structurally and visually much less similar to the approved refined mockups than the backend/business behavior is to the intent. The complaint about UI difference is supported by source evidence. At the same time, expecting the full enterprise shell/design-system outcome from W1 is inconsistent with the original intent, which explicitly deferred those items.", fill=AMBER_LIGHT, accent=AMBER)

add_page_break(doc)
add_heading(doc, "Why the difference occurred", 1)
causes = [
    ("Expectation boundary", "The original intent explicitly deferred full shell/auth and full design-system migration. Some stakeholder expectations were therefore outside W1."),
    ("No authoritative Figma source", "The approved UX was text/ASCII plus HTML concepts; it did not provide immutable component/frame geometry for exact construction."),
    ("Design rules were not enforced", "The refined mapping required @erp/ui tokens/primitives and no local hardcoded visual system, but the Booking app uses independent CSS."),
    ("Acceptance focused more strongly on behavior", "Tests and reviews covered canonical data, APIs, security, retries, events, and service health more strongly than visual composition."),
    ("No blocking visual regression baseline", "There is no W1 implementation screenshot set or Figma-frame comparison that could fail the build."),
    ("Completion and release semantics diverged", "All stages were marked complete while deployment, nginx smoke, performance, and SLO evidence remained blocked."),
]
for title, body in causes:
    p = doc.add_paragraph()
    apply_num(p, decimal_num)
    r = p.add_run(f"{title}. ")
    set_run_font(r, bold=True, color=NAVY)
    set_run_font(p.add_run(body))

add_heading(doc, "Recommended operating model", 1)
table = doc.add_table(rows=1, cols=5)
for cell, text in zip(table.rows[0].cells, ["Gate", "Owner", "Required evidence", "Blocking rule", "AI role"]):
    cell.text = text
gate_rows = [
    ("1. Intent and journey", "Product", "Approved actor, flow, scope, DoD, non-goals", "No design/code without business acceptance", "Draft, challenge, trace"),
    ("2. Design authority", "Design + Product", "Versioned Figma library and feature frames for all states/viewports", "No UI construction from prose alone", "Generate variants and spec summaries"),
    ("3. Story/design trace", "Product + QA", "Stories reference frame/component IDs and behavior", "Missing reference blocks ready status", "Generate AC and trace matrix"),
    ("4. Implementation", "Engineering", "Approved components/tokens; no local visual forks", "Lint blocks local hex/unapproved components", "Build and self-review"),
    ("5. Visual and accessibility", "Design + QA", "Screenshot regression, semantic checks, keyboard/a11y, design signoff", "Mismatch blocks merge", "Run tests and produce evidence"),
    ("6. Live release", "Engineering + Ops", "Fresh PASS manifest, smoke, performance, audit review", "BLOCKED cannot be called complete", "Automate execution and report"),
]
for item in gate_rows:
    cells = table.add_row().cells
    for cell, text in zip(cells, item):
        cell.text = text
style_table(table, [1500, 1450, 2700, 2200, 1510], font_size=7.8)

add_heading(doc, "Concrete policy changes for this repository", 2)
for text in [
    "Change intent status semantics: Completed means every required acceptance gate passed; use Blocked or Conditionally Complete when live gates are not run.",
    "For UI-heavy intents, require a Figma URL/frame manifest or another versioned high-fidelity design artifact at the refined-mockups gate.",
    "Require each UI story and test to reference screen, state, breakpoint, and component identifiers.",
    "Add Booking visual regression baselines for list, create, draft detail, manual pricing, priced, confirm dialog, pending journey, active journey, error, and mobile views.",
    "Add lint rules that reject Booking-local hardcoded colors/fonts and unapproved primitives when the design-system mapping says @erp/ui is mandatory.",
    "Make product/design approval independent and named. An AI review may advise, but cannot be the final UX acceptance authority.",
    "Keep AI-DLC's drift and operational artifacts; they are valuable precisely because they show that the release is not proven.",
]:
    add_bullet(doc, text, bullet_num)

add_heading(doc, "Manager-ready decision", 1)
add_callout(doc, "Decision", "Continue using AI-DLC v2 for enterprise engineering work, with a controlled pilot and explicit design/release governance. Do not represent it as guaranteeing exact output from requirements alone.", fill=TEAL_LIGHT, accent=TEAL)

table = doc.add_table(rows=1, cols=3)
for cell, text in zip(table.rows[0].cells, ["Use AI-DLC when", "Add another authority when", "Do not accept"]):
    cell.text = text
decision_rows = [
    ("Cross-module traceability, contracts, brownfield change, NFRs, deployment, and audit evidence matter", "Visual/UI fidelity matters: Figma design system, feature prototype, design review, visual regression", "A stage marked complete while live acceptance is blocked"),
    ("The outcome can be expressed as observable acceptance criteria", "Business decisions are unresolved: product/CTO decision owner", "Text mockups presented as proof of exact visual output"),
    ("Humans will review the important decisions and gates", "High-risk compliance/security: independent specialists", "AI self-review as the only acceptance authority"),
]
for item in decision_rows:
    cells = table.add_row().cells
    for cell, text in zip(cells, item):
        cell.text = text
style_table(table, [3120, 3120, 3120], font_size=8.5)

add_heading(doc, "Suggested next pilot", 2)
p = doc.add_paragraph("Use the Booking Quote-to-Cash UI gap itself as a short controlled pilot. Freeze approved Figma frames for the list, create, and detail states; migrate the Booking app to @erp/ui; add screenshot regression and design signoff; rerun the live acceptance harness; then measure cycle time, rework, defect count, and fidelity score. This directly tests whether the combined Figma + AI-DLC model addresses the CTO's concern.")

add_page_break(doc)
add_heading(doc, "Appendix A - Key source files", 1)
sources = [
    "docs/examples/booking-quote-to-cash/intent-statement.md",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/ideation/intent-capture/intent-statement.md",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/ideation/rough-mockups/wireframes.md",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/ideation/rough-mockups/user-flow.md",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/ideation/rough-mockups/rough-mockups-questions.md",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/inception/requirements-analysis/requirements.md",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/inception/user-stories/stories.md",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/inception/refined-mockups/mockups.md",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/inception/refined-mockups/interaction-spec.md",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/inception/refined-mockups/design-system-mapping.md",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/inception/refined-mockups/accessibility-checklist.md",
    "apps/booking/app/booking.css",
    "apps/booking/app/bookings/page.tsx",
    "apps/booking/app/bookings/new/BookingCreateForm.tsx",
    "apps/booking/app/bookings/[bookingId]/page.tsx",
    "apps/booking/app/bookings/[bookingId]/BookingValidationPanel.tsx",
    "apps/booking/app/bookings/[bookingId]/JourneyStatusPanel.tsx",
    "packages/ui/src/index.tsx",
    "packages/ui/src/styles.ts",
    "packages/ui/src/primitives.tsx",
    "packages/ui/src/interactive.tsx",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/construction/booking-draft-skeleton/code-generation/code-summary.md",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/construction/build-and-test/build-test-results.md",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/operation/deployment-execution/deployment-log.md",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/operation/performance-validation/load-test-results.md",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/operation/feedback-optimization/drift-report.md",
    "aidlc/spaces/default/intents/260714-booking-quote-cash/aidlc-state.md",
]
for source in sources:
    add_source_line(doc, source)

add_heading(doc, "Appendix B - Evidence limitations", 1)
for text in [
    "The review inspected committed AI-DLC artifacts, current Booking source, package UI source, tests, and operational records in D:\\TST_Codex_integ.",
    "The in-app browser had no available browser binding, so no fresh interactive screenshot or pixel comparison was produced.",
    "No committed W1 implementation screenshot was found under the W1 intent or live-evidence directories.",
    "The 53/100 score is a transparent weighted judgment from source and evidence, not a standardized benchmark or pixel-diff measurement.",
    "The worktree contains ongoing changes for later intents; the report treats the current checked-out Booking source as the implemented state visible in this workspace and does not modify it.",
]:
    add_bullet(doc, text, bullet_num)

# Final document properties and save.
doc.core_properties.title = "AI-DLC v2 Assessment - Booking Quote-to-Cash"
doc.core_properties.subject = "Pros, cons, exact outputs, UI fidelity, and recommended Figma + AI-DLC operating model"
doc.core_properties.author = "OpenAI Codex"
doc.core_properties.keywords = "AI-DLC v2, Booking Quote-to-Cash, Figma, UI fidelity, software delivery"
doc.core_properties.comments = "Generated from repository evidence on 19 July 2026."

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
